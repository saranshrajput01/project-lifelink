"""Build the ESE submission report: copies Cloud_Computing_ESE_LIFELINK.docx and
inserts SDG Alignment + Cloud Services Used sections before the Conclusion.
Run from repo root:  python3 docs/make_ese_report.py
"""

import copy
from pathlib import Path

import docx
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyBboxPatch  # noqa: E402
from docx.shared import Inches  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "Cloud_Computing_ESE_LIFELINK.docx"
OUT = ROOT / "TeamNo_LIFELINK_Report.docx"
DIAGRAM = ROOT / "docs" / "lifelink_cloud_architecture.png"


# ─── Architecture diagram ────────────────────────────────────────────────
def draw_diagram():
    fig, ax = plt.subplots(figsize=(11, 7.2))
    ax.set_xlim(0, 110)
    ax.set_ylim(0, 72)
    ax.axis("off")

    def box(x, y, w, h, title, sub, color):
        ax.add_patch(
            FancyBboxPatch(
                (x, y), w, h, boxstyle="round,pad=0.6", fc=color, ec="#333", lw=1.2
            )
        )
        ax.text(x + w / 2, y + h - 2.6, title, ha="center", va="center",
                fontsize=10.5, weight="bold")
        ax.text(x + w / 2, y + (h - 4) / 2, sub, ha="center", va="center",
                fontsize=8.3, linespacing=1.35)

    def arrow(x1, y1, x2, y2, label="", both=False):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="<|-|>" if both else "-|>",
                                    color="#444", lw=1.3))
        if label:
            ax.text((x1 + x2) / 2 + 0.8, (y1 + y2) / 2, label, fontsize=7.8,
                    color="#222", style="italic")

    box(40, 63, 30, 7, "Users", "Patients · Doctors · Admins (any browser)", "#FFF2CC")
    box(36, 49, 38, 9, "Vercel (Frontend Hosting)",
        "React + TypeScript SPA\nGlobal edge CDN · automatic HTTPS", "#DAE8FC")
    box(31, 29, 48, 14, "Render (Backend PaaS)",
        "FastAPI in Docker · Singapore region\nJWT auth + RBAC · ML risk models\n"
        "RAG chat · reminder scheduler\nHealth check /api/v1/health", "#D5E8D4")
    box(2, 6, 30, 13, "Supabase PostgreSQL",
        "PostgreSQL 17 + pgvector\nusers, records, timeline,\nembeddings, risk scores", "#E1D5E7")
    box(40, 6, 30, 13, "Supabase Storage",
        "S3-compatible private bucket\nlifelink-documents\npre-signed URLs", "#E1D5E7")
    box(78, 6, 30, 13, "External cloud APIs",
        "Groq LLM API (Qwen)\ndocument extraction + chat\nGmail SMTP (reminders)", "#F8CECC")
    box(84, 49, 24, 13, "GitHub",
        "Source repo\nGitHub Actions CI\n(lint · test · build)\nauto-deploy on push", "#F5F5F5")

    arrow(55, 63, 55, 58.6, "HTTPS", both=True)
    arrow(55, 49, 55, 43.6, "HTTPS REST + JWT", both=True)
    arrow(40, 29, 17, 19.6, "SQL (asyncpg / TLS)", both=True)
    arrow(55, 29, 55, 19.6, "boto3 S3 API", both=True)
    arrow(70, 29, 93, 19.6, "HTTPS", both=True)
    arrow(84, 54, 74.6, 54)
    ax.text(79.3, 55.5, "deploy", ha="center", fontsize=7.8, style="italic")
    arrow(90, 49, 79.6, 39, "deploy")

    ax.text(55, 1, "Fig. 1 — LIFELINK live cloud architecture (free tier)",
            ha="center", fontsize=9.5, style="italic")
    fig.tight_layout()
    fig.savefig(DIAGRAM, dpi=200)
    plt.close(fig)


# ─── docx helpers ────────────────────────────────────────────────────────
d = docx.Document(SRC)
P = d.paragraphs
H1, H2, BODY, BULLET = P[81], P[70], P[71], P[28]
TOC_ITEM = P[19]
CONCLUSION = P[81]
TPL = {2: d.tables[0], 3: d.tables[2], 4: d.tables[4]}


def set_text(par_el, text, bold_prefix=None):
    """Replace a paragraph element's text, keeping the first run's formatting."""
    from docx.text.paragraph import Paragraph

    p = Paragraph(par_el, None)
    runs = p.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if bold_prefix:
        runs[0].text = bold_prefix
        runs[0].bold = True
        r2 = copy.deepcopy(runs[0]._r)
        runs[0]._r.addnext(r2)
        from docx.text.run import Run

        rr = Run(r2, p)
        rr.text = text
        rr.bold = False
    else:
        runs[0].text = text
    return par_el


def clone(par, text, **kw):
    return set_text(copy.deepcopy(par._p), text, **kw)


def set_cell(cell, text):
    paras = cell.paragraphs
    for extra in paras[1:]:
        extra._p.getparent().remove(extra._p)
    set_text(paras[0]._p, text)


def make_table(rows):
    ncols = len(rows[0])
    tpl = TPL[ncols]
    tbl = copy.deepcopy(tpl._tbl)
    trs = tbl.findall(docx.oxml.ns.qn("w:tr"))
    header_tr, data_tr = trs[0], trs[1]
    for tr in trs:
        tbl.remove(tr)
    from docx.table import _Cell

    for i, row in enumerate(rows):
        tr = copy.deepcopy(header_tr if i == 0 else data_tr)
        tbl.append(tr)
        for tc, val in zip(tr.findall(docx.oxml.ns.qn("w:tc")), row):
            set_cell(_Cell(tc, None), val)
    return tbl


def image_par():
    p = d.add_paragraph()
    p.alignment = 1
    p.add_run().add_picture(str(DIAGRAM), width=Inches(6.3))
    el = p._p
    el.getparent().remove(el)
    return el


# ─── Content ─────────────────────────────────────────────────────────────
def build_sections():
    els = []
    h1 = lambda t: els.append(clone(H1, t))  # noqa: E731
    h2 = lambda t: els.append(clone(H2, t))  # noqa: E731
    body = lambda t: els.append(clone(BODY, t))  # noqa: E731
    bullet = lambda t, b=None: els.append(clone(BULLET, t, bold_prefix=b))  # noqa: E731
    table = lambda rows: els.append(make_table(rows))  # noqa: E731

    # ── 8. SDG Alignment ──
    h1("8. Alignment with UN Sustainable Development Goals")
    body(
        "LIFELINK is a healthcare product, so its main contribution is to SDG 3 "
        "(Good Health and Well-being). Because it is delivered entirely over the "
        "cloud, it also supports SDG 9 (Industry, Innovation and Infrastructure). "
        "The specific targets it addresses are listed below."
    )
    h2("8.1 SDGs and Targets Supported")
    table([
        ["SDG / Target", "Target (summary)", "How LIFELINK Contributes"],
        ["SDG 3 — Target 3.4",
         "Reduce premature deaths from non-communicable diseases by one third "
         "through prevention and treatment.",
         "ML models flag risk of cardiovascular disease, diabetes and chronic kidney "
         "disease early, with the factors behind each score and lifestyle "
         "recommendations. Medicine and follow-up reminders support adherence."],
        ["SDG 3 — Target 3.8",
         "Achieve universal health coverage, including access to quality "
         "essential health services.",
         "Any patient with a browser can store and carry a complete medical history "
         "for free, so a new doctor or emergency team sees the full picture."],
        ["SDG 3 — Target 3.d",
         "Strengthen capacity for early warning, risk reduction and management "
         "of health risks.",
         "Abnormal lab values, a health score and risk levels are surfaced on the "
         "dashboard automatically after each upload."],
        ["SDG 9 — Target 9.c",
         "Significantly increase access to information and communications "
         "technology.",
         "AI document reading, an AI health assistant and risk prediction are "
         "offered as a web service that needs no installation or special hardware."],
    ])
    h2("8.2 How the Project Contributes")
    bullet(" Scattered paper and PDF records become one searchable, chronological "
           "timeline, so medical history is not lost between hospitals.",
           "Unified records:")
    bullet(" OCR and an LLM pull out medications, diagnoses and lab values, so "
           "patients can understand their own reports.", "Health literacy:")
    bullet(" Risk scores for three common chronic diseases encourage patients to "
           "seek care before complications appear.", "Prevention:")
    bullet(" Family profiles let one person manage records for elderly parents or "
           "children.", "Family care:")
    h2("8.3 Expected Impact")
    table([
        ["Dimension", "Expected Impact"],
        ["Social",
         "Better continuity of care, fewer medication conflicts, and earlier "
         "detection of chronic disease. Patients and caregivers gain control over "
         "their own health data."],
        ["Economic",
         "Fewer repeated diagnostic tests because previous reports are always "
         "available. Early risk detection can avoid costlier late-stage treatment. "
         "The platform itself runs at zero hosting cost on free cloud tiers, so it "
         "can be offered free to users."],
        ["Environmental",
         "Digital records replace printed reports and repeat tests, reducing paper "
         "and lab consumables. Shared, multi-tenant cloud infrastructure uses "
         "hardware more efficiently than running a dedicated server per clinic."],
    ])

    # ── 9. Cloud Services Used ──
    h1("9. Cloud Services Used")
    body(
        "LIFELINK is live on the internet using free-tier cloud platforms. "
        "Frontend: https://project-lifelink-1.vercel.app — Backend API: "
        "https://lifelink-backend-tunc.onrender.com. The AWS design from Sections "
        "3–6 is written as Terraform code (infra/ folder) and is the scale-up path "
        "once the project needs production capacity."
    )
    h2("9.1 Services in the Live Deployment")
    table([
        ["#", "Cloud Service", "Model", "Used For"],
        ["1", "Vercel", "PaaS + CDN",
         "Hosts the React frontend on a global edge network with automatic HTTPS; "
         "rebuilds on every push to GitHub."],
        ["2", "Render Web Service", "PaaS",
         "Runs the FastAPI backend as a Docker container (Singapore region), with a "
         "health check on /api/v1/health and auto-deploy from GitHub."],
        ["3", "Supabase PostgreSQL", "DBaaS",
         "Managed PostgreSQL 17 with the pgvector extension; stores users, records, "
         "timeline, risk assessments, reminders and RAG embeddings."],
        ["4", "Supabase Storage", "Object storage (STaaS)",
         "Private S3-compatible bucket (lifelink-documents) for uploaded reports; "
         "accessed via boto3 with time-limited pre-signed download URLs."],
        ["5", "Groq Cloud API", "AI API (SaaS)",
         "LLM inference (Qwen model) for document data extraction, summaries and "
         "the AI health assistant, with multi-key rotation for quota."],
        ["6", "GitHub + GitHub Actions", "DevOps SaaS",
         "Source control and CI pipeline (lint, test, build); pushes to main "
         "trigger cloud redeployment."],
        ["7", "Gmail SMTP", "SaaS",
         "Sends medicine and follow-up reminder emails from the background "
         "scheduler."],
    ])
    body(
        "AWS services prepared in Terraform for the scale-up path: Amazon VPC "
        "(public/private subnets, NAT Gateway), Application Load Balancer, Amazon "
        "ECS on Fargate, Amazon RDS for PostgreSQL 16, Amazon S3 (encryption, "
        "versioning, public access blocked), Amazon ECR, AWS Secrets Manager, AWS "
        "IAM roles and Amazon CloudWatch Logs."
    )
    h2("9.2 Architecture Diagram")
    els.append(image_par())
    body(
        "A user opens the site from Vercel's CDN. The React app calls the FastAPI "
        "backend on Render over HTTPS with a JWT token. The backend reads and writes "
        "structured data in Supabase PostgreSQL, stores files in Supabase Storage, "
        "and calls the Groq API for AI processing. GitHub pushes trigger CI and "
        "redeploy both the frontend and the backend."
    )
    h2("9.3 Benefits of Using the Cloud")
    table([
        ["Benefit", "How LIFELINK Gets It"],
        ["Scalability",
         "Vercel's CDN serves any number of frontend users. The backend is a "
         "stateless container, so Render (or ECS Fargate auto-scaling on AWS) can run "
         "more instances behind a load balancer. Storage and database grow on demand."],
        ["Cost",
         "Pay-as-you-go with no hardware purchase. The current deployment costs $0 "
         "per month on free tiers; paid tiers are only needed when traffic grows."],
        ["Availability",
         "Managed providers handle uptime, patching and backups. Render restarts the "
         "container if the health check fails. Supabase manages database backups. "
         "Note: the free Render tier sleeps after 15 minutes idle, so the first "
         "request takes about a minute."],
        ["Security",
         "HTTPS everywhere, JWT authentication with role-based access, bcrypt "
         "password hashing, a private storage bucket with expiring pre-signed URLs, "
         "and secrets kept in platform environment variables instead of code."],
        ["Accessibility",
         "Patients and doctors use it from any device with a browser, anywhere, with "
         "no installation. Records are available during emergencies or when changing "
         "doctors."],
    ])
    return els


# ─── Assemble ────────────────────────────────────────────────────────────
def main():
    draw_diagram()
    anchor = CONCLUSION._p
    for el in build_sections():
        anchor.addprevious(el)

    # Renumber Conclusion / References headings and TOC
    set_text(P[81]._p, "10. Conclusion")
    set_text(P[86]._p, "11. References")
    toc8 = P[19]
    toc8._p.addprevious(clone(TOC_ITEM, "8.  Alignment with UN Sustainable Development Goals"))
    toc8._p.addprevious(clone(TOC_ITEM, "9.  Cloud Services Used"))
    set_text(P[19]._p, "10.  Conclusion")
    set_text(P[20]._p, "11.  References")

    # Cover page
    P[6]._p.addprevious(clone(P[6], "Team No: ____"))
    set_text(P[9]._p, "Submission Date: October 2026")

    # Conclusion addition
    P[85]._p.addnext(clone(P[85],
        "The project is already live on free-tier cloud platforms (Vercel, Render "
        "and Supabase), which shows that the design works in practice at zero cost. "
        "It directly supports UN SDG 3 by helping patients detect chronic disease "
        "risk early and keep their medical history in one accessible place."))

    # References
    last = P[-1]
    refs = [
        '[9] United Nations, "Goal 3: Ensure healthy lives and promote well-being for '
        'all at all ages," UN DESA. Available: https://sdgs.un.org/goals/goal3',
        '[10] United Nations, "Goal 9: Industry, Innovation and Infrastructure," UN '
        'DESA. Available: https://sdgs.un.org/goals/goal9',
        '[11] Supabase, "Supabase Documentation — Database, pgvector and Storage." '
        'Available: https://supabase.com/docs',
        '[12] Render, "Render Documentation — Web Services and Blueprints." '
        'Available: https://render.com/docs',
        '[13] Vercel, "Vercel Documentation." Available: https://vercel.com/docs',
    ]
    for r in reversed(refs):
        last._p.addnext(clone(last, r))

    d.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()
